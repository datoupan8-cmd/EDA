"""Build the Chinese Component V4 implementation report as a styled DOCX."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Component_V4_融合方案与实现说明.md"
OUTPUT = ROOT / "Component_V4_融合方案与实现说明.docx"
ACCENT = "334155"
LIGHT = "F1F5F9"
INK = RGBColor(31, 41, 55)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", 9.5, color=RGBColor(22, 78, 99))
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True, color=INK)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, color=INK)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color in (
        ("Title", 25, RGBColor(0, 0, 0)),
        ("Heading 1", 17, RGBColor(0, 0, 0)),
        ("Heading 2", 13, RGBColor(0, 0, 0)),
        ("Heading 3", 11.5, RGBColor(0, 0, 0)),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.keep_with_next = True
        p_pr = style._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("PCB 原理图自动解析  Component V4.1")
    set_run_font(run, size=8.5, color=RGBColor(0, 0, 0))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("本地非官方诊断报告  |  ")
    set_run_font(run, size=8, color=RGBColor(100, 116, 139))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(78)
    p_pr = p._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)
    run = p.add_run("Component V4.1\n融合方案与实现说明")
    set_run_font(run, size=26, bold=True, color=RGBColor(0, 0, 0))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("YOLO Symbol Detection + OCR/Text Understanding\n+ Component–Text Association")
    set_run_font(run, size=13, color=RGBColor(0, 0, 0))

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(9.0)
    facts = (
        ("开发数据", "最新版官方 0001～0150"),
        ("封存数据", "0151～0200 + 10 Golden，未使用"),
        ("默认主线", "BEST + RapidOCR 1.4.4 全图 + EasyOCR 全分块"),
        ("评测性质", "本地非官方诊断，OFFICIAL_SCORE = FALSE"),
    )
    for row, (label, value) in zip(table.rows, facts):
        set_cell_shading(row.cells[0], LIGHT)
        for cell, text in zip(row.cells, (label, value)):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=10, bold=(cell is row.cells[0]), color=INK)
    doc.add_page_break()


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F6FA")
    p_pr.append(shd)
    run = p.add_run("\n".join(lines))
    set_run_font(run, "Consolas", 8.6, color=RGBColor(30, 64, 89))


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    parsed = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(parsed) < 2:
        return
    rows = [parsed[0]] + parsed[2:]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        row_pr = row._tr.get_or_add_trPr()
        row_pr.append(OxmlElement("w:cantSplit"))
        for c_idx, value in enumerate(values):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=70, bottom=70)
            if r_idx == 0:
                set_cell_shading(cell, ACCENT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, value)
            for run in p.runs:
                set_run_font(run, size=8.3, bold=(r_idx == 0), color=RGBColor(255, 255, 255) if r_idx == 0 else INK)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_markdown_body(doc: Document, text: str) -> None:
    lines = text.splitlines()
    # The DOCX title page replaces the Markdown H1.
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|---"):
            table_lines = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("## "):
            if line.startswith("## 七、"):
                doc.add_page_break()
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1


def build() -> Path:
    doc = Document()
    setup_document(doc)
    add_title_page(doc)
    add_markdown_body(doc, SOURCE.read_text(encoding="utf-8-sig"))
    for section in doc.sections:
        section.page_height = Cm(27.94)
        section.page_width = Cm(21.59)
    doc.core_properties.title = "Component V4.1 融合方案与实现说明"
    doc.core_properties.subject = "PCB 原理图自动解析 Component 前端融合"
    doc.core_properties.author = "PCB Competition Team"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
